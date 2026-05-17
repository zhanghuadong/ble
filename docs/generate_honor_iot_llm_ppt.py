from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


OUTPUT = "docs/honor_iot_llm_solution.pptx"


class Theme:
    navy = RGBColor(18, 32, 59)
    blue = RGBColor(35, 99, 235)
    cyan = RGBColor(14, 165, 233)
    green = RGBColor(16, 185, 129)
    amber = RGBColor(245, 158, 11)
    red = RGBColor(239, 68, 68)
    purple = RGBColor(124, 58, 237)
    slate = RGBColor(71, 85, 105)
    light = RGBColor(248, 250, 252)
    line = RGBColor(203, 213, 225)
    white = RGBColor(255, 255, 255)
    dark = RGBColor(15, 23, 42)


FONT = "Microsoft YaHei"


def set_run(run, size=18, bold=False, color=Theme.dark):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_text(frame, text, size=18, bold=False, color=Theme.dark, align=PP_ALIGN.LEFT):
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)


def add_bg(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = Theme.light


def add_title(slide, title, subtitle=None):
    add_bg(slide)
    box = slide.shapes.add_textbox(Cm(0.8), Cm(0.45), Cm(26.5), Cm(1.1))
    set_text(box.text_frame, title, 28, True, Theme.navy)
    if subtitle:
        sub = slide.shapes.add_textbox(Cm(0.85), Cm(1.55), Cm(24.5), Cm(0.65))
        set_text(sub.text_frame, subtitle, 11, False, Theme.slate)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Cm(0.85), Cm(2.18), Cm(2.0), Cm(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = Theme.blue
    line.line.fill.background()


def add_footer(slide, idx):
    box = slide.shapes.add_textbox(Cm(24.6), Cm(14.45), Cm(2.2), Cm(0.35))
    set_text(box.text_frame, f"{idx:02d}", 8, False, Theme.slate, PP_ALIGN.RIGHT)


def rounded_box(slide, x, y, w, h, text, fill, line=None, font=15, bold=True, color=Theme.white):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    shp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(shp.text_frame, text, font, bold, color, PP_ALIGN.CENTER)
    return shp


def card(slide, x, y, w, h, title, body, accent=Theme.blue):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = Theme.white
    shp.line.color.rgb = Theme.line
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Cm(0.12), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    tx = slide.shapes.add_textbox(x + Cm(0.35), y + Cm(0.25), w - Cm(0.55), h - Cm(0.4))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 15, True, Theme.navy)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(8)
    r2 = p2.add_run()
    r2.text = body
    set_run(r2, 10, False, Theme.slate)
    return shp


def arrow(slide, x1, y1, x2, y2, color=Theme.slate, width=2.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.line.end_arrowhead = True
    return conn


def add_bullets(slide, x, y, w, h, bullets, size=14, color=Theme.dark):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.text = text
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def add_icon(slide, kind, x, y, color):
    if kind == "phone":
        body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Cm(1.25), Cm(2.1))
        body.fill.solid()
        body.fill.fore_color.rgb = Theme.white
        body.line.color.rgb = color
        screen = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + Cm(0.18), y + Cm(0.28), Cm(0.9), Cm(1.35))
        screen.fill.solid()
        screen.fill.fore_color.rgb = color
        screen.line.fill.background()
        btn = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(0.52), y + Cm(1.72), Cm(0.22), Cm(0.22))
        btn.fill.solid()
        btn.fill.fore_color.rgb = color
        btn.line.fill.background()
    elif kind == "hub":
        body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + Cm(0.45), Cm(1.7), Cm(1.25))
        body.fill.solid()
        body.fill.fore_color.rgb = color
        body.line.fill.background()
        for i in range(3):
            led = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(0.28 + i * 0.42), y + Cm(1.02), Cm(0.18), Cm(0.18))
            led.fill.solid()
            led.fill.fore_color.rgb = Theme.white
            led.line.fill.background()
    elif kind == "server":
        for i in range(3):
            rack = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + Cm(i * 0.48), Cm(1.7), Cm(0.38))
            rack.fill.solid()
            rack.fill.fore_color.rgb = color
            rack.line.fill.background()
            dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(1.35), y + Cm(i * 0.48 + 0.11), Cm(0.13), Cm(0.13))
            dot.fill.solid()
            dot.fill.fore_color.rgb = Theme.white
            dot.line.fill.background()
    elif kind == "chip":
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + Cm(0.2), y + Cm(0.2), Cm(1.3), Cm(1.3))
        chip.fill.solid()
        chip.fill.fore_color.rgb = color
        chip.line.fill.background()
        core = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + Cm(0.52), y + Cm(0.52), Cm(0.66), Cm(0.66))
        core.fill.solid()
        core.fill.fore_color.rgb = Theme.white
        core.line.fill.background()
        for i in range(4):
            left_pin = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                x + Cm(0.02),
                y + Cm(0.35 + i * 0.28),
                Cm(0.2),
                Cm(0.08),
            )
            left_pin.fill.solid()
            left_pin.fill.fore_color.rgb = color
            left_pin.line.fill.background()
            right_pin = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                x + Cm(1.48),
                y + Cm(0.35 + i * 0.28),
                Cm(0.2),
                Cm(0.08),
            )
            right_pin.fill.solid()
            right_pin.fill.fore_color.rgb = color
            right_pin.line.fill.background()
    elif kind == "bulb":
        bulb = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(0.25), y, Cm(1.2), Cm(1.2))
        bulb.fill.solid()
        bulb.fill.fore_color.rgb = color
        bulb.line.fill.background()
        base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + Cm(0.58), y + Cm(1.1), Cm(0.54), Cm(0.42))
        base.fill.solid()
        base.fill.fore_color.rgb = Theme.slate
        base.line.fill.background()
    elif kind == "llm":
        cloud = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CLOUD, x, y, Cm(1.9), Cm(1.2))
        cloud.fill.solid()
        cloud.fill.fore_color.rgb = color
        cloud.line.fill.background()
        tx = slide.shapes.add_textbox(x + Cm(0.38), y + Cm(0.35), Cm(1.2), Cm(0.35))
        set_text(tx.text_frame, "LLM", 9, True, Theme.white, PP_ALIGN.CENTER)


def add_role_card(slide, x, title, body, icon, color):
    y = Cm(4.2)
    add_icon(slide, icon, x + Cm(1.15), y - Cm(1.55), color)
    card(slide, x, y, Cm(4.6), Cm(3.15), title, body, color)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    blob = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ARC, Cm(18.4), Cm(-1.5), Cm(9), Cm(9))
    blob.line.color.rgb = Theme.cyan
    blob.line.width = Pt(16)
    box = slide.shapes.add_textbox(Cm(1.0), Cm(1.0), Cm(15.8), Cm(2.4))
    set_text(box.text_frame, "荣耀智慧空间 App\n授权中枢与 LLM 控制 IoT 方案", 30, True, Theme.navy)
    sub = slide.shapes.add_textbox(Cm(1.08), Cm(3.55), Cm(16.0), Cm(0.7))
    set_text(sub.text_frame, "HONOR Connect SDK · Node MCP Server · ESP32-S3 · 本地 IoT 设备", 15, False, Theme.slate)

    xs = [Cm(2.0), Cm(6.6), Cm(11.1), Cm(15.6), Cm(20.1)]
    labels = ["荣耀智慧空间", "中枢 / SDK", "MCP Server", "ESP32-S3", "IoT 设备"]
    icons = ["phone", "hub", "server", "chip", "bulb"]
    colors = [Theme.blue, Theme.purple, Theme.cyan, Theme.green, Theme.amber]
    for i, (x, label, icon, color) in enumerate(zip(xs, labels, icons, colors)):
        add_icon(slide, icon, x + Cm(0.65), Cm(7.1), color)
        rounded_box(slide, x, Cm(9.25), Cm(3.1), Cm(0.88), label, color, font=11)
        if i < len(xs) - 1:
            arrow(slide, x + Cm(3.2), Cm(9.7), xs[i + 1] - Cm(0.15), Cm(9.7), Theme.slate)

    llm = rounded_box(slide, Cm(10.9), Cm(5.05), Cm(4.2), Cm(1.0), "大语言模型入口", Theme.navy, font=14)
    add_icon(slide, "llm", Cm(9.25), Cm(4.95), Theme.navy)
    arrow(slide, Cm(13.0), Cm(6.08), Cm(12.7), Cm(9.22), Theme.navy, 2.5)
    note = slide.shapes.add_textbox(Cm(1.05), Cm(13.2), Cm(20), Cm(0.5))
    set_text(note.text_frame, "目标：让荣耀智慧空间完成用户授权与入口控制，让中枢统一管理 IoT，并让 LLM 通过受控工具安全操作设备。", 12, False, Theme.slate)


def slide_goals(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "1. 需求拆解：三个入口，一个控制核心", "把荣耀 App、IoT 控制和大语言模型能力收敛到同一套中枢控制面")
    add_bullets(
        slide,
        Cm(1.1),
        Cm(3.0),
        Cm(9.5),
        Cm(4.8),
        [
            "荣耀智慧空间 App：负责用户侧授权、绑定、控制入口和场景入口",
            "中枢：负责权限、设备模型、命令总线、协议转换和状态同步",
            "LLM：通过 MCP / Function Calling 调用中枢工具，不直接触碰硬件",
            "IoT 设备：由 ESP32-S3 或其他本地总线执行最终动作",
        ],
        15,
    )
    rounded_box(slide, Cm(13.0), Cm(3.0), Cm(4.0), Cm(1.1), "授权入口", Theme.blue)
    rounded_box(slide, Cm(13.0), Cm(6.1), Cm(4.0), Cm(1.1), "自然语言入口", Theme.purple)
    rounded_box(slide, Cm(13.0), Cm(9.2), Cm(4.0), Cm(1.1), "自动化入口", Theme.green)
    rounded_box(slide, Cm(20.0), Cm(6.1), Cm(4.2), Cm(1.2), "中枢控制核心", Theme.navy)
    for y in [Cm(3.55), Cm(6.65), Cm(9.75)]:
        arrow(slide, Cm(17.0), y, Cm(19.9), Cm(6.7), Theme.slate)
    rounded_box(slide, Cm(20.2), Cm(9.4), Cm(3.8), Cm(0.95), "IoT 执行层", Theme.amber, font=13)
    arrow(slide, Cm(22.1), Cm(7.3), Cm(22.1), Cm(9.35), Theme.slate)
    add_footer(slide, idx)


def slide_architecture(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "2. 推荐总体架构", "所有控制请求统一进入中枢，避免权限、日志和状态分裂")
    y = Cm(6.2)
    nodes = [
        ("荣耀智慧空间 App", "phone", Theme.blue),
        ("中枢\nHONOR Connect SDK", "hub", Theme.purple),
        ("Node MCP Server\n工具与协议适配", "server", Theme.cyan),
        ("ESP32-S3\n硬件执行节点", "chip", Theme.green),
        ("IoT 设备\n灯/窗帘/传感器", "bulb", Theme.amber),
    ]
    xs = [Cm(0.9), Cm(5.95), Cm(11.2), Cm(16.45), Cm(21.2)]
    for i, ((label, icon, color), x) in enumerate(zip(nodes, xs)):
        add_icon(slide, icon, x + Cm(1.05), Cm(4.0), color)
        rounded_box(slide, x, y, Cm(4.25), Cm(1.35), label, color, font=12)
        if i < len(nodes) - 1:
            arrow(slide, x + Cm(4.3), y + Cm(0.68), xs[i + 1] - Cm(0.1), y + Cm(0.68), Theme.slate)
    rounded_box(slide, Cm(10.9), Cm(2.8), Cm(4.8), Cm(1.1), "大语言模型", Theme.navy)
    arrow(slide, Cm(13.3), Cm(3.9), Cm(13.3), Cm(6.15), Theme.navy)
    card(slide, Cm(1.1), Cm(9.1), Cm(7.2), Cm(2.2), "关键原则", "荣耀 App 与 LLM 都是入口；真正的权限校验、设备状态和命令执行必须在中枢统一处理。", Theme.blue)
    card(slide, Cm(9.8), Cm(9.1), Cm(7.2), Cm(2.2), "协议建议", "中枢到 Node 可用 HTTP/WebSocket/JSON-RPC；Node 到 ESP32-S3 优先 MQTT，其次 WebSocket/HTTP/RS485。", Theme.cyan)
    card(slide, Cm(18.4), Cm(9.1), Cm(7.0), Cm(2.2), "扩展方向", "后续可增加 Zigbee、BLE、Matter、Modbus 网关，仍通过同一设备模型暴露给荣耀与 LLM。", Theme.green)
    add_footer(slide, idx)


def slide_roles(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "3. 系统角色分工", "每一层只承担自己最擅长的职责")
    add_role_card(slide, Cm(0.9), "荣耀智慧空间 App", "用户授权、家庭/房间入口、设备控制 UI、场景触发。", "phone", Theme.blue)
    add_role_card(slide, Cm(6.0), "中枢", "HONOR SDK 适配、权限中心、设备模型、命令总线、状态同步。", "hub", Theme.purple)
    add_role_card(slide, Cm(11.1), "Node MCP Server", "向 LLM 暴露工具，完成意图到设备命令的受控转换。", "server", Theme.cyan)
    add_role_card(slide, Cm(16.2), "ESP32-S3", "GPIO/PWM/传感器/继电器/RS485 等底层硬件执行。", "chip", Theme.green)
    add_role_card(slide, Cm(21.3), "IoT 设备", "最终被控制或采集数据的灯、窗帘、插座、传感器等。", "bulb", Theme.amber)
    add_footer(slide, idx)


def slide_honor_flow(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "4. 链路 A：荣耀智慧空间 App 控制", "用户在 App 中授权/绑定中枢后，通过中枢下发 IoT 控制")
    steps = [
        ("1", "用户点击\n打开客厅灯", Theme.blue),
        ("2", "HONOR SDK\n转发控制指令", Theme.purple),
        ("3", "中枢校验\n用户与设备权限", Theme.navy),
        ("4", "协议转换\n发送给执行节点", Theme.cyan),
        ("5", "ESP32-S3\n控制继电器/调光", Theme.green),
        ("6", "状态回传\n同步 App 展示", Theme.amber),
    ]
    x = Cm(1.0)
    for i, (num, text, color) in enumerate(steps):
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(0.92), Cm(4.0), Cm(0.8), Cm(0.8))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.fill.background()
        set_text(circ.text_frame, num, 14, True, Theme.white, PP_ALIGN.CENTER)
        rounded_box(slide, x, Cm(5.1), Cm(2.65), Cm(1.4), text, color, font=10)
        if i < len(steps) - 1:
            arrow(slide, x + Cm(2.72), Cm(5.8), x + Cm(3.42), Cm(5.8), Theme.slate)
        x += Cm(4.05)
    card(slide, Cm(2.0), Cm(9.0), Cm(10.3), Cm(2.7), "请求示例", "{\"deviceId\":\"living_room_light\", \"property\":\"power\", \"value\":true, \"source\":\"honor_app\"}", Theme.blue)
    card(slide, Cm(14.4), Cm(9.0), Cm(10.3), Cm(2.7), "结果要求", "每个命令必须有 cmdId、ACK、状态回写与审计日志，确保 App 展示和设备真实状态一致。", Theme.green)
    add_footer(slide, idx)


def slide_llm_flow(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "5. 链路 B：大语言模型通过 MCP 控制", "LLM 调用的是安全工具，不是 MQTT、GPIO 或任意代码执行")
    bubble = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(1.1), Cm(3.0), Cm(7.0), Cm(1.6))
    bubble.fill.solid()
    bubble.fill.fore_color.rgb = Theme.white
    bubble.line.color.rgb = Theme.line
    set_text(bubble.text_frame, "“把客厅灯调暗一点，窗帘关到一半”", 14, True, Theme.navy, PP_ALIGN.CENTER)
    add_icon(slide, "llm", Cm(3.5), Cm(5.0), Theme.purple)
    rounded_box(slide, Cm(10.0), Cm(3.25), Cm(5.7), Cm(1.1), "MCP 工具调用", Theme.purple)
    arrow(slide, Cm(8.15), Cm(3.8), Cm(9.95), Cm(3.8), Theme.purple)
    tool_texts = ["list_devices()", "get_device_state()", "set_device_property()", "run_scene()"]
    for i, text in enumerate(tool_texts):
        rounded_box(slide, Cm(10.4), Cm(5.0 + i * 1.1), Cm(4.8), Cm(0.72), text, Theme.white, Theme.line, 10, False, Theme.dark)
    rounded_box(slide, Cm(18.0), Cm(5.4), Cm(5.4), Cm(1.1), "中枢权限与安全策略", Theme.navy)
    arrow(slide, Cm(15.8), Cm(6.65), Cm(17.95), Cm(5.95), Theme.slate)
    rounded_box(slide, Cm(18.25), Cm(8.0), Cm(4.9), Cm(1.05), "IoT 设备执行", Theme.green)
    arrow(slide, Cm(20.7), Cm(6.5), Cm(20.7), Cm(7.95), Theme.slate)
    card(slide, Cm(1.1), Cm(10.7), Cm(7.3), Cm(2.15), "安全边界", "LLM 只能操作中枢返回的设备与白名单工具；高风险动作必须二次确认。", Theme.red)
    card(slide, Cm(9.6), Cm(10.7), Cm(7.3), Cm(2.15), "用户权限继承", "LLM 不拥有超级管理员权限，只继承当前用户授权给中枢的可控范围。", Theme.purple)
    card(slide, Cm(18.1), Cm(10.7), Cm(7.3), Cm(2.15), "状态一致性", "执行后由中枢更新状态缓存，并同步给荣耀 App 与对话侧。", Theme.green)
    add_footer(slide, idx)


def slide_hub_core(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "6. 中枢核心职责", "中枢不只是转发器，而是权限、模型、命令和状态的统一控制面")
    rounded_box(slide, Cm(10.5), Cm(5.55), Cm(5.3), Cm(1.25), "统一命令总线", Theme.navy)
    items = [
        (Cm(1.2), Cm(3.2), "设备模型", "设备、房间、能力、当前状态统一抽象", Theme.blue),
        (Cm(17.9), Cm(3.2), "权限管理", "用户、家庭、房间、设备可控范围", Theme.purple),
        (Cm(1.2), Cm(9.0), "协议适配", "HONOR / MCP / MQTT / BLE / RS485", Theme.cyan),
        (Cm(17.9), Cm(9.0), "审计与安全", "cmdId、ACK、日志、二次确认、风控", Theme.red),
    ]
    for x, y, title, body, color in items:
        card(slide, x, y, Cm(7.0), Cm(2.35), title, body, color)
        arrow(slide, x + Cm(7.0 if x < Cm(10) else 0), y + Cm(1.15), Cm(13.15), Cm(6.15), Theme.slate)
    add_footer(slide, idx)


def slide_merge(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "7. ESP32-S3 执行层能否与中枢合并？", "可以合并，但关键取决于中枢硬件能否承载 HONOR SDK、Node/MCP 和权限系统")
    rounded_box(slide, Cm(2.0), Cm(3.2), Cm(8.0), Cm(1.0), "方案 A：Android / Linux 中枢 + ESP32-S3", Theme.green, font=13)
    rounded_box(slide, Cm(2.0), Cm(4.8), Cm(3.1), Cm(0.9), "荣耀 SDK", Theme.purple, font=11)
    rounded_box(slide, Cm(5.6), Cm(4.8), Cm(3.1), Cm(0.9), "MCP/LLM", Theme.cyan, font=11)
    rounded_box(slide, Cm(9.2), Cm(4.8), Cm(3.1), Cm(0.9), "ESP 执行", Theme.green, font=11)
    arrow(slide, Cm(5.1), Cm(5.25), Cm(5.55), Cm(5.25), Theme.slate)
    arrow(slide, Cm(8.7), Cm(5.25), Cm(9.15), Cm(5.25), Theme.slate)
    add_bullets(slide, Cm(2.3), Cm(6.35), Cm(9.5), Cm(2.6), ["推荐用于产品化与多设备扩展", "职责清晰，后续可接多个 ESP32/总线节点"], 12)

    rounded_box(slide, Cm(14.8), Cm(3.2), Cm(8.0), Cm(1.0), "方案 B：ESP32-S3 小型中枢 + 执行器", Theme.amber, font=13)
    rounded_box(slide, Cm(15.0), Cm(4.8), Cm(3.2), Cm(0.9), "本地 HTTP/MQTT", Theme.cyan, font=10)
    rounded_box(slide, Cm(18.5), Cm(4.8), Cm(3.2), Cm(0.9), "GPIO/PWM", Theme.green, font=10)
    arrow(slide, Cm(18.25), Cm(5.25), Cm(18.45), Cm(5.25), Theme.slate)
    add_bullets(slide, Cm(15.0), Cm(6.35), Cm(9.5), Cm(3.4), ["适合原型或少量继电器/传感器", "通常难以直接承载完整 HONOR Connect SDK", "不适合复杂 LLM、数据库与多用户权限"], 12)

    rounded_box(slide, Cm(5.0), Cm(11.5), Cm(16.5), Cm(1.0), "结论：要接荣耀授权与 LLM，优先保留 Android/Linux 中枢；ESP32-S3 作为执行节点更稳。", Theme.navy, font=13)
    add_footer(slide, idx)


def slide_deploy(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "8. 三种部署形态选择", "根据产品目标选择合并或分层，不必一开始做重架构")
    headers = ["部署形态", "适用场景", "优点", "限制"]
    rows = [
        ["Android/Linux 中枢 + ESP32 节点", "可扩展产品、多房间、多协议", "荣耀 SDK、Node MCP、权限和执行职责清晰", "硬件数量更多"],
        ["Android/Linux 中枢直接控设备", "网关有 GPIO/RS485/Zigbee/BLE", "减少 ESP32 层级，部署简单", "对中枢硬件接口要求高"],
        ["ESP32-S3 小型中枢", "单设备原型、低成本验证", "硬件最少，执行链路短", "难承载完整荣耀 SDK 与复杂 LLM 控制"],
    ]
    x0, y0 = Cm(0.9), Cm(3.1)
    widths = [Cm(5.6), Cm(6.2), Cm(7.1), Cm(6.2)]
    for i, h in enumerate(headers):
        rounded_box(slide, x0 + sum(widths[:i]), y0, widths[i] - Cm(0.1), Cm(0.85), h, Theme.navy, font=11)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            fill = RGBColor(255, 255, 255) if r % 2 == 0 else RGBColor(241, 245, 249)
            shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x0 + sum(widths[:c]), y0 + Cm(1.05 + r * 2.25), widths[c] - Cm(0.1), Cm(2.0))
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
            shp.line.color.rgb = Theme.line
            shp.text_frame.word_wrap = True
            set_text(shp.text_frame, text, 10, c == 0, Theme.dark, PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT)
    add_footer(slide, idx)


def slide_security(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "9. 授权与安全设计", "让 LLM 控制设备前，先建立用户授权、工具白名单和风险策略")
    add_icon(slide, "phone", Cm(1.6), Cm(3.0), Theme.blue)
    rounded_box(slide, Cm(3.5), Cm(3.4), Cm(4.0), Cm(0.9), "用户授权中枢", Theme.blue, font=12)
    arrow(slide, Cm(7.6), Cm(3.85), Cm(9.8), Cm(3.85), Theme.slate)
    rounded_box(slide, Cm(9.9), Cm(3.4), Cm(4.0), Cm(0.9), "生成本地 Token", Theme.purple, font=12)
    arrow(slide, Cm(14.0), Cm(3.85), Cm(16.2), Cm(3.85), Theme.slate)
    rounded_box(slide, Cm(16.3), Cm(3.4), Cm(4.2), Cm(0.9), "LLM 继承用户权限", Theme.navy, font=12)
    arrow(slide, Cm(18.4), Cm(4.35), Cm(18.4), Cm(6.0), Theme.slate)
    rounded_box(slide, Cm(15.8), Cm(6.05), Cm(5.2), Cm(1.0), "中枢策略校验", Theme.red, font=13)

    card(slide, Cm(1.1), Cm(7.5), Cm(5.8), Cm(2.5), "设备密钥", "每台 ESP32-S3 使用唯一 deviceId 与 secret，MQTT/HTTP 都要认证。", Theme.green)
    card(slide, Cm(7.5), Cm(7.5), Cm(5.8), Cm(2.5), "工具白名单", "LLM 只允许调用 list/get/set/run_scene 等预定义工具。", Theme.purple)
    card(slide, Cm(13.9), Cm(7.5), Cm(5.8), Cm(2.5), "高危二次确认", "门锁、燃气阀、大功率插座等操作必须确认。", Theme.red)
    card(slide, Cm(20.3), Cm(7.5), Cm(5.1), Cm(2.5), "审计日志", "记录 source、userId、cmdId、deviceId、result。", Theme.blue)
    rounded_box(slide, Cm(4.2), Cm(12.0), Cm(17.9), Cm(0.95), "安全底线：大语言模型永远不应绕过中枢直接控制硬件。", Theme.navy, font=13)
    add_footer(slide, idx)


def slide_mvp(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "10. MVP 到产品化路线", "先打通最小闭环，再逐步扩展设备、协议和自动化能力")
    phases = [
        ("阶段 1\n最小闭环", "荣耀绑定中枢\n控制一个灯/继电器\n状态回传 App 与 LLM", Theme.blue),
        ("阶段 2\n多设备管理", "设备注册/发现\nMQTT Topic 规范\n在线离线与 OTA", Theme.cyan),
        ("阶段 3\n场景与自动化", "房间/分组/场景\n本地规则引擎\n传感器联动", Theme.green),
        ("阶段 4\n安全与运维", "审计日志\n高危确认\n诊断与故障恢复", Theme.purple),
    ]
    y = Cm(5.2)
    prev_x = None
    for i, (title, body, color) in enumerate(phases):
        x = Cm(1.1 + i * 6.25)
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(1.9), Cm(3.15), Cm(1.0), Cm(1.0))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.fill.background()
        set_text(circ.text_frame, str(i + 1), 14, True, Theme.white, PP_ALIGN.CENTER)
        card(slide, x, y, Cm(5.3), Cm(3.2), title, body, color)
        if prev_x is not None:
            arrow(slide, prev_x + Cm(5.35), Cm(6.8), x - Cm(0.05), Cm(6.8), Theme.slate)
        prev_x = x
    rounded_box(slide, Cm(3.4), Cm(11.0), Cm(19.8), Cm(1.0), "MVP 验收标准：荣耀 App 可授权并控制设备；LLM 可通过 MCP 查询/控制；设备状态可一致回传。", Theme.navy, font=13)
    add_footer(slide, idx)


def slide_protocols(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "11. 协议与数据模型建议", "使用统一设备模型屏蔽荣耀、LLM 与底层设备协议差异")
    card(slide, Cm(1.0), Cm(3.0), Cm(7.4), Cm(3.0), "统一设备模型", "id / type / room / capabilities / state\n示例：light.power、brightness、colorTemperature", Theme.blue)
    card(slide, Cm(9.6), Cm(3.0), Cm(7.4), Cm(3.0), "中枢 ↔ Node", "优先 WebSocket + JSON-RPC；REST 适合配置与查询；事件使用订阅推送。", Theme.cyan)
    card(slide, Cm(18.2), Cm(3.0), Cm(7.4), Cm(3.0), "Node ↔ ESP32-S3", "优先 MQTT；Topic：iot/{deviceId}/command、state、event、ota。", Theme.green)

    rounded_box(slide, Cm(3.1), Cm(8.2), Cm(4.8), Cm(0.9), "HONOR 模型", Theme.purple, font=12)
    rounded_box(slide, Cm(10.8), Cm(8.2), Cm(4.8), Cm(0.9), "中枢统一模型", Theme.navy, font=12)
    rounded_box(slide, Cm(18.5), Cm(8.2), Cm(4.8), Cm(0.9), "设备私有协议", Theme.amber, font=12)
    arrow(slide, Cm(7.95), Cm(8.65), Cm(10.75), Cm(8.65), Theme.slate)
    arrow(slide, Cm(15.65), Cm(8.65), Cm(18.45), Cm(8.65), Theme.slate)
    add_bullets(
        slide,
        Cm(4.0),
        Cm(10.5),
        Cm(18.5),
        Cm(2.6),
        [
            "荣耀侧能力受品类与 SDK 限制，需要映射到标准设备能力。",
            "LLM 侧只暴露人类可理解的工具与设备名称，最终仍转换成统一模型。",
            "底层设备协议可替换，不影响荣耀 App 与 LLM 的上层体验。",
        ],
        13,
    )
    add_footer(slide, idx)


def slide_risks(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "12. 风险与待确认事项", "这些点决定方案边界与实现复杂度")
    risks = [
        ("HONOR Connect SDK 能力边界", "是否允许第三方中枢暴露自定义设备；是否需要认证、品类白名单或生态审核。", Theme.red),
        ("设备模型映射限制", "荣耀支持的标准品类和属性可能无法覆盖所有自定义 IoT 能力。", Theme.amber),
        ("本地发现与绑定流程", "中枢如何被 App 发现，Node 如何发现 ESP32-S3，设备如何归属到房间。", Theme.blue),
        ("断网与离线控制", "是否要求局域网闭环可用；云端 LLM 不可用时是否保留本地场景。", Theme.green),
        ("高风险设备策略", "门锁、燃气阀、安防等必须明确二次确认、日志和回滚策略。", Theme.purple),
        ("运维与 OTA", "ESP32-S3 固件签名、版本回滚、设备故障诊断和远程日志。", Theme.cyan),
    ]
    for i, (title, body, color) in enumerate(risks):
        x = Cm(1.0 + (i % 2) * 12.7)
        y = Cm(3.0 + (i // 2) * 3.25)
        card(slide, x, y, Cm(11.4), Cm(2.45), title, body, color)
    add_footer(slide, idx)


def slide_conclusion(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    box = slide.shapes.add_textbox(Cm(1.1), Cm(1.0), Cm(24.5), Cm(1.4))
    set_text(box.text_frame, "推荐结论", 32, True, Theme.navy)
    rounded_box(slide, Cm(1.2), Cm(3.1), Cm(23.8), Cm(1.25), "荣耀 App 和 LLM 都是入口；中枢必须是唯一可信控制核心。", Theme.navy, font=18)
    cards = [
        ("产品化推荐", "Android/Linux 中枢集成 HONOR SDK 与 Node MCP；ESP32-S3 作为可扩展执行节点。", Theme.green),
        ("原型验证", "可以让 ESP32-S3 做小型中枢 + 执行器，但荣耀 SDK 与 LLM 能力通常需要代理或上位机。", Theme.amber),
        ("安全原则", "LLM 只能通过 MCP 工具进入中枢，继承用户授权，经过白名单、风控和审计。", Theme.red),
    ]
    for i, (title, body, color) in enumerate(cards):
        card(slide, Cm(1.2 + i * 8.1), Cm(6.0), Cm(7.3), Cm(3.1), title, body, color)
    rounded_box(slide, Cm(4.3), Cm(11.4), Cm(17.5), Cm(1.1), "下一步：先做“一个灯”的 MVP，验证授权、控制、状态回传与 LLM 工具调用闭环。", Theme.blue, font=14)
    add_footer(slide, idx)


def build():
    prs = Presentation()
    prs.slide_width = Cm(26.6667)
    prs.slide_height = Cm(15.0)
    slide_cover(prs)
    slide_goals(prs, 2)
    slide_architecture(prs, 3)
    slide_roles(prs, 4)
    slide_honor_flow(prs, 5)
    slide_llm_flow(prs, 6)
    slide_hub_core(prs, 7)
    slide_merge(prs, 8)
    slide_deploy(prs, 9)
    slide_security(prs, 10)
    slide_mvp(prs, 11)
    slide_protocols(prs, 12)
    slide_risks(prs, 13)
    slide_conclusion(prs, 14)
    prs.save(OUTPUT)
    print(f"Generated {OUTPUT}")


if __name__ == "__main__":
    build()
